"""
CertiGenius — Import Participants Page
Upload CSV or Excel files to create certificate batches.
"""
import os
import uuid
from datetime import datetime, timezone

import streamlit as st

from models import CertificateBatch, Participant, TemplateGroup, Base
from utils.db import get_session, init_connection
from utils.excel_parser import parse_file

st.set_page_config(page_title="Import Participants — CertiGenius", page_icon="👥", layout="wide")

# ---- Init DB ----
engine = init_connection()
Base.metadata.create_all(engine)

st.title("👥 Import Participants")
st.markdown("Upload a CSV or Excel file with participant data to create a new batch.")

# ---- Upload Form ----
with st.form("import_form", clear_on_submit=True):
    col1, col2, col3 = st.columns(3)
    with col1:
        batch_name = st.text_input(
            "Batch Name",
            placeholder=f"e.g., Summer Workshop {datetime.now().year}",
        )
    with col2:
        event_type = st.selectbox(
            "Event Type",
            options=["ISA Event", "Non-ISA Event"],
            help="Select whether this batch is for an ISA or Non-ISA event. Templates will be chosen accordingly.",
        )
        event_type_value = "isa" if event_type == "ISA Event" else "non_isa"
    with col3:
        session = get_session()
        try:
            template_groups = session.query(TemplateGroup).filter_by(is_deleted=False).all()
            tg_options = {g.name: g.id for g in template_groups}
        finally:
            session.close()

        selected_group = st.selectbox(
            "Template Group (optional)",
            options=["— No template group —"] + list(tg_options.keys()),
        )
        template_group_id = tg_options.get(selected_group) if selected_group != "— No template group —" else None

    uploaded_file = st.file_uploader(
        "Select CSV or Excel File",
        type=["csv", "xlsx", "xls"],
        help="Required columns: Name, Email. Optional: Prize Position + any custom fields.",
    )

    submitted = st.form_submit_button("🚀 Import & Create Batch", use_container_width=True, type="primary")

    if submitted:
        if not uploaded_file:
            st.error("Please select a file to upload.")
            st.stop()

        # Save the uploaded file temporarily
        upload_dir = os.path.join(os.getcwd(), "uploads", "participants")
        os.makedirs(upload_dir, exist_ok=True)
        
        ext = os.path.splitext(uploaded_file.name)[1].lower()
        temp_filename = f"{uuid.uuid4().hex}{ext}"
        temp_path = os.path.join(upload_dir, temp_filename)
        
        with open(temp_path, "wb") as f:
            f.write(uploaded_file.getbuffer())

        # Parse the file
        parse_result = parse_file(temp_path)

        if parse_result.errors:
            for err in parse_result.errors[:10]:
                st.error(err)

        if parse_result.valid_count == 0:
            st.error("No valid participants found in the file.")
            st.stop()

        # Create batch and participants
        bname = batch_name.strip() or f"Batch {datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M')}"
        
        session = get_session()
        try:
            batch = CertificateBatch(
                name=bname,
                template_group_id=template_group_id,
                event_type=event_type_value,
                status="pending",
                total_count=0,
            )
            session.add(batch)
            session.flush()

            for parsed in parse_result.participants:
                if not parsed.is_valid():
                    continue
                participant = Participant(
                    name=parsed.name,
                    email=parsed.email,
                    prize_position=parsed.prize_position,
                    batch_id=batch.id,
                )
                participant.set_extra_data(parsed.extra_data)
                session.add(participant)
                batch.total_count += 1

            session.commit()
            st.success(f"✅ Created batch '{bname}' with {parse_result.valid_count} participants!")

            # Store the batch ID to navigate after rerun
            st.session_state["last_batch_id"] = batch.id

        except Exception as e:
            session.rollback()
            st.error(f"Error creating batch: {e}")
        finally:
            session.close()

# Handle navigation after form submission (outside the form)
if "last_batch_id" in st.session_state:
    batch_id = st.session_state.pop("last_batch_id")
    if st.button(f"🔍 View Batch #{batch_id}", use_container_width=True, type="primary"):
        st.switch_page(f"pages/4_Batch_Detail.py")
        st.query_params["batch_id"] = str(batch_id)

# ---- Sample Download Link ----
st.markdown("---")
st.subheader("📎 Need a sample file?")

col1, col2 = st.columns(2)
with col1:
    if st.button("📥 Download Sample Excel", use_container_width=True):
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill

        wb = Workbook()
        ws = wb.active
        ws.title = "Participants"
        headers = ["Name", "Email", "Prize Position", "Event", "Date"]
        ws.append(headers)
        sample_data = [
            ["John Doe", "john@example.com", "1st", "Annual Conference 2024", "2024-12-01"],
            ["Jane Smith", "jane@example.com", "2nd", "Annual Conference 2024", "2024-12-01"],
            ["Bob Johnson", "bob@example.com", "Participation", "Annual Conference 2024", "2024-12-01"],
        ]
        for row in sample_data:
            ws.append(row)

        header_fill = PatternFill(start_color="667EEA", end_color="667EEA", fill_type="solid")
        header_font = Font(color="FFFFFF", bold=True, size=11)
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font

        for col_cells in ws.columns:
            max_length = max((len(str(c.value)) for c in col_cells if c.value), default=0)
            ws.column_dimensions[col_cells[0].column_letter].width = max_length + 4

        import io
        output = io.BytesIO()
        wb.save(output)
        output.seek(0)

        st.download_button(
            label="📥 Download sample_participants.xlsx",
            data=output,
            file_name="sample_participants.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            use_container_width=True,
        )

with col2:
    st.info(
        "The file should have at least **Name** and **Email** columns. "
        "**Prize Position** is optional (defaults to 'Participation'). "
        "Any extra columns become placeholder fields for your templates."
    )

# ---- Recent Batches ----
st.markdown("---")
st.subheader("📦 Recent Batches")

session = get_session()
try:
    recent = (
        session.query(CertificateBatch)
        .order_by(CertificateBatch.created_at.desc())
        .limit(5)
        .all()
    )
finally:
    session.close()

if recent:
    rows = []
    for b in recent:
        rows.append({
            "ID": b.id,
            "Name": b.name,
            "Status": b.status,
            "Participants": b.total_count,
            "Created": b.created_at.strftime("%b %d, %Y") if b.created_at else "—",
        })

    st.dataframe(rows, use_container_width=True, hide_index=True)
else:
    st.caption("No batches yet.")
